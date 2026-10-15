#!/usr/bin/env python3
"""ローカル請求書の対象月設定と、再計算済みコピーの値固定。"""
import argparse
import calendar
import json
import math
import os
import posixpath
import re
import tempfile
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from xml.dom import minidom
from zipfile import ZipFile

from openpyxl import load_workbook
from openpyxl.utils.datetime import to_excel

MAIN = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'
REL = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'


class InvalidInvoice(ValueError):
    pass


def children(node, name):
    return [n for n in node.childNodes if n.nodeType == n.ELEMENT_NODE and n.localName == name]


def text_value(node):
    return ''.join(n.data for n in node.childNodes if n.nodeType == n.TEXT_NODE)


def remove_children(node, names):
    for child in list(node.childNodes):
        if child.nodeType == child.ELEMENT_NODE and child.localName in names:
            node.removeChild(child)


def element(parent, name, text=None):
    prefix = parent.prefix + ':' if parent.prefix else ''
    result = parent.ownerDocument.createElementNS(MAIN, prefix + name)
    if text is not None:
        result.appendChild(parent.ownerDocument.createTextNode(text))
    parent.appendChild(result)
    return result


def sheet_parts(archive):
    doc = minidom.parseString(archive.read('xl/workbook.xml'))
    rels = minidom.parseString(archive.read('xl/_rels/workbook.xml.rels'))
    targets = {r.getAttribute('Id'): r.getAttribute('Target') for r in rels.documentElement.childNodes if r.nodeType == r.ELEMENT_NODE}
    result = {}
    for sheet in doc.getElementsByTagNameNS(MAIN, 'sheet'):
        target = targets[sheet.getAttributeNS(REL, 'id')]
        path = target.lstrip('/') if target.startswith('/') else posixpath.normpath('xl/' + target)
        result[sheet.getAttribute('name')] = path
    return result


def read_book(path):
    if path.suffix.lower() != '.xlsx':
        raise InvalidInvoice('入力は .xlsx が必要です')
    with ZipFile(path) as archive:
        if any(n.startswith('xl/externalLinks/') or n.endswith('vbaProject.bin') for n in archive.namelist()):
            raise InvalidInvoice('外部リンク・マクロを含むブックは対象外です')
        for name in archive.namelist():
            if name.startswith('xl/tables/') and name.endswith('.xml'):
                table = minidom.parseString(archive.read(name))
                if any(table.getElementsByTagNameNS(MAIN, tag) for tag in ('calculatedColumnFormula', 'totalsRowFormula')):
                    raise InvalidInvoice('テーブル数式を含むブックは対象外です')
    book = load_workbook(path, data_only=False)
    for sheet in book:
        for row in sheet:
            for cell in row:
                if cell.data_type == 'f' and not isinstance(cell.value, str):
                    raise InvalidInvoice('配列数式・データテーブルは対象外です')
    return book


def target_month(value):
    if not re.fullmatch(r'\d{4}-\d{2}', value):
        raise InvalidInvoice('対象月は YYYY-MM で指定してください')
    return date.fromisoformat(value + '-01')


def configured_cell(book, ref):
    sheet, address = ref['sheet'], ref['cell']
    if sheet not in book.sheetnames or not re.fullmatch(r'[A-Z]{1,3}[1-9]\d*', address):
        raise InvalidInvoice('profile のシート名・単一セル参照を確認してください')
    return book[sheet][address]


def write_package(source, output, transform):
    # 検証失敗や途中終了で既存出力を壊さず、既存ファイルは上書きしない。
    if output.exists():
        raise InvalidInvoice('出力先が既に存在します')
    with tempfile.NamedTemporaryFile(dir=output.parent, suffix='.xlsx', delete=False) as tmp:
        temporary = Path(tmp.name)
    try:
        with ZipFile(source) as src, ZipFile(temporary, 'w') as dst:
            for item in src.infolist():
                data = transform(item.filename, src.read(item.filename))
                if data is not None:
                    dst.writestr(item, data)
        # hard link の作成は既存出力へ置換しないため、同時実行でも上書きしない。
        os.link(temporary, output)
    finally:
        temporary.unlink(missing_ok=True)


def prepare(source, output, profile, month):
    book = read_book(source)
    cell = configured_cell(book, profile['month_cell'])
    if not cell.is_date or cell.data_type == 'f':
        raise InvalidInvoice('month_cell は日付書式の入力セルが必要です')
    with ZipFile(source) as archive:
        parts = sheet_parts(archive)
    month_part = parts[cell.parent.title]
    sheet_paths = set(parts.values())

    def transform(name, data):
        if name not in sheet_paths and name != 'xl/workbook.xml':
            return data
        doc = minidom.parseString(data)
        if name == 'xl/workbook.xml':
            calc = doc.getElementsByTagNameNS(MAIN, 'calcPr')
            calc = calc[0] if calc else element(doc.documentElement, 'calcPr')
            for key, value in {'calcMode': 'auto', 'fullCalcOnLoad': '1', 'forceFullCalc': '1'}.items():
                calc.setAttribute(key, value)
        else:
            for node in doc.getElementsByTagNameNS(MAIN, 'c'):
                if children(node, 'f'):
                    remove_children(node, {'v'})
                if name == month_part and node.getAttribute('r') == cell.coordinate:
                    remove_children(node, {'v', 'is'})
                    node.setAttribute('t', 'n')
                    element(node, 'v', str(to_excel(month, book.epoch)))
        return doc.toxml(encoding='utf-8')

    write_package(source, output, transform)
    return {'stage': 'prepared', 'recalculation_required': True}


def signature(book):
    return {sheet.title: {cell.coordinate: (cell.data_type, cell.value) for row in sheet for cell in row if cell.value is not None} for sheet in book}


def defined_names(path):
    with ZipFile(path) as archive:
        doc = minidom.parseString(archive.read('xl/workbook.xml'))
    return sorted((tuple(sorted(node.attributes.items())), text_value(node)) for node in doc.getElementsByTagNameNS(MAIN, 'definedName'))


def print_settings(sheet):
    return tuple(str(value) for value in (sheet.print_area, sheet.page_setup, sheet.page_margins, sheet.print_options, sheet.print_title_rows, sheet.print_title_cols, sheet.sheet_properties.pageSetUpPr, sheet.row_breaks, sheet.col_breaks))


def inspect(book, profile, month):
    month_value = configured_cell(book, profile['month_cell']).value
    if isinstance(month_value, datetime):
        month_value = month_value.date()
    if month_value != month:
        raise InvalidInvoice('対象月セルが指定月の初日と一致しません')
    for ref in profile['date_checks']:
        offset = ref.get('month_offset', 0)
        index = month.year * 12 + month.month - 1 + offset
        year, m = divmod(index, 12)
        m += 1
        day = ref['day']
        if day not in ('first', 'last'):
            raise InvalidInvoice('date_checks.day は first または last が必要です')
        expected = date(year, m, 1 if day == 'first' else calendar.monthrange(year, m)[1])
        value = configured_cell(book, ref).value
        if isinstance(value, datetime):
            value = value.date()
        if value != expected:
            raise InvalidInvoice('日付検査に失敗しました: ' + ref['cell'])
    if not profile['amount_cells'] or not profile['print_sheets']:
        raise InvalidInvoice('金額セルと印刷対象シートの指定が必要です')
    for ref in profile['amount_cells']:
        value = configured_cell(book, ref).value
        if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
            raise InvalidInvoice('金額が有限の数値ではありません: ' + ref['cell'])
        if 'expected' in ref and Decimal(str(value)) != Decimal(str(ref['expected'])):
            raise InvalidInvoice('期待金額と一致しません: ' + ref['cell'])
    for name in profile['print_sheets']:
        sheet = book[name]
        if sheet.sheet_state != 'visible' or not sheet.print_area or not sheet.page_setup.paperSize or not sheet.page_setup.orientation:
            raise InvalidInvoice('表示状態・印刷範囲・用紙・向きを確認してください')


def export(prepared, recalculated, output, profile, month):
    before, after = read_book(prepared), read_book(recalculated)
    if signature(before) != signature(after) or defined_names(prepared) != defined_names(recalculated):
        raise InvalidInvoice('再計算前後で入力値・数式が変わっています')
    for sheet in before:
        if print_settings(sheet) != print_settings(after[sheet.title]):
            raise InvalidInvoice('再計算前後で印刷設定が変わっています')
    with ZipFile(prepared) as archive:
        parts = sheet_parts(archive)
        for path in parts.values():
            doc = minidom.parseString(archive.read(path))
            if any(children(node, 'v') for node in doc.getElementsByTagNameNS(MAIN, 'c') if children(node, 'f')):
                raise InvalidInvoice('prepare で作成した未計算ブックを指定してください')
    with ZipFile(recalculated) as archive:
        paths = set(sheet_parts(archive).values())
    formula_count = 0

    def transform(name, data):
        nonlocal formula_count
        if name == 'xl/calcChain.xml':
            return None
        if name in ('xl/_rels/workbook.xml.rels', '[Content_Types].xml'):
            doc = minidom.parseString(data)
            removed = False
            for node in list(doc.documentElement.childNodes):
                if node.nodeType != node.ELEMENT_NODE:
                    continue
                if node.getAttribute('Type').endswith('/calcChain') or node.getAttribute('PartName') == '/xl/calcChain.xml':
                    doc.documentElement.removeChild(node)
                    removed = True
            return doc.toxml(encoding='utf-8') if removed else data
        if name not in paths:
            return data
        doc = minidom.parseString(data)
        for node in doc.getElementsByTagNameNS(MAIN, 'c'):
            if node.getAttribute('t') == 'e':
                raise InvalidInvoice('Excel エラーを含むセルがあります')
            if not children(node, 'f'):
                continue
            values = children(node, 'v')
            kind = node.getAttribute('t')
            if not values or (kind != 'str' and not text_value(values[0]).strip()):
                raise InvalidInvoice('数式の計算キャッシュがありません: ' + node.getAttribute('r'))
            value = text_value(values[0])
            if kind in ('', 'n') and not math.isfinite(float(value)):
                raise InvalidInvoice('数式の計算結果が有限ではありません')
            if kind == 'b' and value not in ('0', '1'):
                raise InvalidInvoice('真偽値キャッシュが不正です')
            if kind not in ('', 'n', 'b', 'str', 'd'):
                raise InvalidInvoice('未対応のキャッシュ形式です')
            remove_children(node, {'f'})
            if kind == 'str':
                remove_children(node, {'v'})
                node.setAttribute('t', 'inlineStr')
                text = element(element(node, 'is'), 't', value)
                text.setAttribute('xml:space', 'preserve')
            formula_count += 1
        return doc.toxml(encoding='utf-8')

    # 一時出力で全検査を済ませてから、最終名を作る。
    with tempfile.TemporaryDirectory(dir=output.parent) as directory:
        candidate = Path(directory) / 'checked.xlsx'
        write_package(recalculated, candidate, transform)
        final = load_workbook(candidate, data_only=True)
        inspect(final, profile, month)
        if any(cell.data_type == 'f' for sheet in read_book(candidate) for row in sheet for cell in row):
            raise InvalidInvoice('数式が残っています')
        os.link(candidate, output)
    return {'stage': 'exported', 'formulas_frozen': formula_count, 'date_checks': len(profile['date_checks']) + 1, 'amount_checks': len(profile['amount_cells']), 'print_settings_checked': True, 'print_preview': 'manual_required'}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest='command', required=True)
    for name in ('prepare', 'export'):
        command = commands.add_parser(name)
        command.add_argument('source', type=Path)
        if name == 'export':
            command.add_argument('recalculated', type=Path)
        command.add_argument('output', type=Path)
        command.add_argument('--profile', type=Path, required=True)
        command.add_argument('--month', required=True)
    args = parser.parse_args()
    try:
        profile = json.loads(args.profile.read_text())
        month = target_month(args.month)
        if args.output.exists():
            raise InvalidInvoice('出力先が既に存在します')
        if args.command == 'prepare':
            report = prepare(args.source, args.output, profile, month)
        else:
            report = export(args.source, args.recalculated, args.output, profile, month)
        print(json.dumps(report, ensure_ascii=False))
    except (InvalidInvoice, ValueError, KeyError, TypeError, OSError, InvalidOperation) as error:
        # 入力の値や数式をエラーログへ展開しない。
        message = str(error) if isinstance(error, InvalidInvoice) else 'ファイル形式・profile・出力先を確認してください'
        parser.exit(1, message + '\n')


if __name__ == '__main__':
    main()
