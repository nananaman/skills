"""合成請求書だけを使った公開 CLI の検証。"""
import json
import subprocess
import sys
import tempfile
import unittest
from datetime import datetime
from pathlib import Path
from zipfile import ZipFile
from xml.etree import ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.table import Table, TableFormula

SCRIPT = Path(__file__).resolve().parents[1] / 'scripts/invoice_xlsx.py'
NS = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}


class InvoiceCliTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.template = self.root / 'template.xlsx'
        wb = Workbook()
        ws = wb.active
        ws.title = 'Invoice'
        ws['B2'] = datetime(2026, 9, 1)
        ws['B2'].number_format = 'yyyy-mm-dd'
        ws['B3'] = '=EOMONTH(B2,0)'
        ws['B3'].number_format = 'yyyy-mm-dd'
        ws['B5'] = 120
        ws['B6'] = '=B5*1.08'
        ws['B7'] = '=0'
        ws.print_area = 'A1:D10'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = 'portrait'
        wb.save(self.template)
        self.profile = self.root / 'profile.json'
        self.profile.write_text(json.dumps({
            'date_inputs': [{'sheet': 'Invoice', 'cell': 'B2', 'value': 'month_start'}],
            'date_checks': [{'sheet': 'Invoice', 'cell': 'B3', 'day': 'last', 'month_offset': 0}],
            'amount_cells': [{'sheet': 'Invoice', 'cell': 'B6'}, {'sheet': 'Invoice', 'cell': 'B7'}],
            'print_sheets': ['Invoice'],
        }))
        self.prepared = self.root / 'prepared.xlsx'
        self.recalculated = self.root / 'recalculated.xlsx'
        self.output = self.root / 'output.xlsx'

    def cli(self, *args):
        return subprocess.run([sys.executable, str(SCRIPT), *map(str, args)], text=True, capture_output=True)

    def prepare(self):
        result = self.cli('prepare', self.template, self.prepared, '--profile', self.profile, '--month', '2026-10')
        self.assertEqual(result.returncode, 0, result.stderr)

    def recalculate(self, caches=None):
        # Excel相当の計算結果を、合成fixtureのOOXMLへ設定する。
        caches = caches or {'B3': '46326', 'B6': '129.6', 'B7': '0'}
        with ZipFile(self.prepared) as src, ZipFile(self.recalculated, 'w') as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == 'xl/worksheets/sheet1.xml':
                    xml = ET.fromstring(data)
                    for cell in xml.findall('.//s:c', NS):
                        if cell.attrib['r'] in caches:
                            value = cell.find('s:v', NS)
                            if value is None:
                                value = ET.SubElement(cell, '{%s}v' % NS['s'])
                            value.text = caches[cell.attrib['r']]
                    data = ET.tostring(xml)
                dst.writestr(item, data)

    def export(self):
        return self.cli('export', self.prepared, self.recalculated, self.output, '--profile', self.profile, '--month', '2026-10')

    def test_preparation_updates_month_and_invalidates_every_formula_cache(self):
        self.prepare()
        wb = load_workbook(self.prepared)
        self.assertEqual(wb['Invoice']['B2'].value, datetime(2026, 10, 1))
        self.assertEqual(wb['Invoice']['B6'].value, '=B5*1.08')
        with ZipFile(self.prepared) as book:
            xml = ET.fromstring(book.read('xl/worksheets/sheet1.xml'))
            self.assertTrue(all(c.find('s:v', NS) is None for c in xml.findall('.//s:c', NS) if c.find('s:f', NS) is not None))
        self.assertEqual(load_workbook(self.template)['Invoice']['B2'].value, datetime(2026, 9, 1))


    def test_export_freezes_zero_and_amounts_and_preserves_package_and_print_settings(self):
        self.prepare()
        self.recalculate()

        result = self.export()

        self.assertEqual(result.returncode, 0, result.stderr)
        report = json.loads(result.stdout)
        self.assertEqual(report['formulas_frozen'], 3)
        self.assertEqual(report['print_preview'], 'manual_required')
        final = load_workbook(self.output)
        self.assertEqual(final['Invoice']['B6'].value, 129.6)
        self.assertEqual(final['Invoice']['B7'].value, 0)
        self.assertFalse(any(c.data_type == 'f' for ws in final for row in ws for c in row))
        with ZipFile(self.recalculated) as src, ZipFile(self.output) as dst:
            for name in src.namelist():
                if not name.startswith('xl/worksheets/'):
                    self.assertEqual(src.read(name), dst.read(name))
        self.assertEqual(str(final['Invoice'].print_area), "'Invoice'!$A$1:$D$10")

    def test_export_rejects_missing_cache_without_creating_output(self):
        self.prepare()
        self.recalculate({'B3': '46326', 'B6': '129.6'})

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('キャッシュがありません', result.stderr)
        self.assertFalse(self.output.exists())

    def test_export_rejects_changed_inputs_and_keeps_existing_output(self):
        self.prepare()
        self.recalculate()
        wb = load_workbook(self.recalculated)
        wb['Invoice']['B5'] = 999
        wb.save(self.recalculated)
        self.output.write_bytes(b'keep existing')

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertEqual(self.output.read_bytes(), b'keep existing')
        self.output.unlink()
        result = self.export()
        self.assertNotEqual(result.returncode, 0)
        self.assertIn('入力値・数式', result.stderr)
        self.assertFalse(self.output.exists())

    def test_export_rejects_wrong_month_end_date(self):
        self.prepare()
        self.recalculate({'B3': '46325', 'B6': '129.6', 'B7': '0'})

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('日付検査', result.stderr)
        self.assertFalse(self.output.exists())

    def test_prepare_rejects_invalid_month(self):
        result = self.cli('prepare', self.template, self.prepared, '--profile', self.profile, '--month', '2026-13')

        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.prepared.exists())

    def test_export_rejects_unexpected_amount_when_operator_sets_expected_value(self):
        profile = json.loads(self.profile.read_text())
        profile['amount_cells'][0]['expected'] = '130'
        self.profile.write_text(json.dumps(profile))
        self.prepare()
        self.recalculate()

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('期待金額', result.stderr)
        self.assertFalse(self.output.exists())


    def rewrite_zip(self, path, change, extras=None):
        temporary = path.with_suffix('.tmp')
        with ZipFile(path) as src, ZipFile(temporary, 'w') as dst:
            for item in src.infolist():
                dst.writestr(item, change(item.filename, src.read(item.filename)))
            for name, data in (extras or {}).items():
                dst.writestr(name, data)
        temporary.replace(path)

    def test_export_removes_calculation_chain_and_its_package_references(self):
        self.prepare()
        self.recalculate()
        def add_chain(name, data):
            if name == '[Content_Types].xml':
                return data.replace(b'</Types>', b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
            if name == 'xl/_rels/workbook.xml.rels':
                return data.replace(b'</Relationships>', b'<Relationship Id="chain" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
            return data
        self.rewrite_zip(self.recalculated, add_chain, {'xl/calcChain.xml': '<calcChain xmlns="%s"><c r="B6" i="1"/></calcChain>' % NS['s']})

        result = self.export()

        self.assertEqual(result.returncode, 0, result.stderr)
        with ZipFile(self.output) as archive:
            self.assertNotIn('xl/calcChain.xml', archive.namelist())
            self.assertNotIn(b'calcChain', archive.read('[Content_Types].xml'))
            self.assertNotIn(b'calcChain', archive.read('xl/_rels/workbook.xml.rels'))

    def test_export_rejects_formula_error_in_hidden_sheet(self):
        wb = load_workbook(self.template)
        hidden = wb.create_sheet('Hidden')
        hidden.sheet_state = 'hidden'
        hidden['A1'] = '=1/0'
        wb.save(self.template)
        self.prepare()
        self.recalculate()
        def add_error(name, data):
            if name == 'xl/worksheets/sheet2.xml':
                xml = ET.fromstring(data)
                node = xml.find('.//s:c', NS)
                node.set('t', 'e')
                ET.SubElement(node, '{%s}v' % NS['s']).text = '#DIV/0!'
                return ET.tostring(xml)
            return data
        self.rewrite_zip(self.recalculated, add_error)

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('Excel エラー', result.stderr)
        self.assertFalse(self.output.exists())

    def test_export_retains_prepared_print_settings_when_recalculated_copy_changes_them(self):
        self.prepare()
        self.recalculate()
        def change_print(name, data):
            if name == 'xl/worksheets/sheet1.xml':
                return data.replace(b'portrait', b'landscape')
            return data
        self.rewrite_zip(self.recalculated, change_print)

        result = self.export()

        self.assertEqual(result.returncode, 0, result.stderr)
        self.assertEqual(load_workbook(self.output)['Invoice'].page_setup.orientation, 'portrait')


    def test_export_rejects_numeric_input_changed_to_boolean(self):
        wb = load_workbook(self.template)
        wb.active['B5'] = 1
        wb.save(self.template)
        self.prepare()
        self.recalculate()
        self.rewrite_zip(self.recalculated, lambda name, data: data.replace(b'r="B5" t="n"', b'r="B5" t="b"') if name == 'xl/worksheets/sheet1.xml' else data)

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())

    def test_export_rejects_changed_defined_name(self):
        wb = load_workbook(self.template)
        wb.defined_names.add(DefinedName('Rate', attr_text='1.08'))
        wb.active['B6'] = '=B5*Rate'
        wb.save(self.template)
        self.prepare()
        self.recalculate()
        self.rewrite_zip(self.recalculated, lambda name, data: data.replace(b'>1.08</', b'>1.10</') if name == 'xl/workbook.xml' else data)

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertFalse(self.output.exists())

    def test_export_retains_prepared_scaling_and_manual_page_breaks(self):
        wb = load_workbook(self.template)
        wb.active.sheet_properties.pageSetUpPr.fitToPage = True
        wb.active.row_breaks.append(Break(id=5))
        wb.save(self.template)
        self.prepare()
        for old, new in [(b'fitToPage="1"', b'fitToPage="0"'), (b'id="5"', b'id="6"')]:
            with self.subTest(setting=old):
                self.output.unlink(missing_ok=True)
                self.recalculate()
                self.rewrite_zip(self.recalculated, lambda name, data: data.replace(old, new) if name == 'xl/worksheets/sheet1.xml' else data)

                result = self.export()

                self.assertEqual(result.returncode, 0, result.stderr)
                final = load_workbook(self.output)['Invoice']
                self.assertTrue(final.sheet_properties.pageSetUpPr.fitToPage)
                self.assertEqual(final.row_breaks.brk[0].id, 5)

    def test_prepare_rejects_table_formula_definitions(self):
        wb = load_workbook(self.template)
        ws = wb.active
        ws['F1'] = 'Amount'
        ws['F2'] = '=B5*1.08'
        table = Table(displayName='InvoiceTable', ref='F1:F2')
        table._initialise_columns()
        table.tableColumns[0].name = 'Amount'
        table.tableColumns[0].calculatedColumnFormula = TableFormula(attr_text='B5*1.08')
        ws.add_table(table)
        wb.save(self.template)

        result = self.cli('prepare', self.template, self.prepared, '--profile', self.profile, '--month', '2026-10')

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('テーブル数式', result.stderr)
        self.assertFalse(self.prepared.exists())


    def volatile_date_template(self, with_print_area):
        # 実情報を含まない、日付数式だけを持つテンプレート構造。
        wb = Workbook()
        ws = wb.active
        ws.title = '請求書 '
        for address, formula in [('G3', '=TODAY()'), ('B7', '=EOMONTH(TODAY(),0)'), ('B8', '=TODAY()')]:
            ws[address] = formula
            ws[address].number_format = 'yyyy-mm-dd'
        ws['G10'] = '=120*1.08'
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.orientation = 'portrait'
        wb.save(self.template)
        profile = {
            'date_inputs': [
                {'sheet': '請求書 ', 'cell': 'G3', 'value': 'issue_date'},
                {'sheet': '請求書 ', 'cell': 'B7', 'value': 'due_date'},
            ],
            'date_checks': [{'sheet': '請求書 ', 'cell': 'B7', 'day': 'last'}],
            'amount_cells': [{'sheet': '請求書 ', 'cell': 'G10', 'expected': '129.6'}],
            'print_sheets': ['請求書 '],
        }
        if with_print_area:
            profile['print_areas'] = {'請求書 ': 'A1:G12'}
        self.profile.write_text(json.dumps(profile))

    def dated_cli(self, command):
        args = [self.template, self.prepared] if command == 'prepare' else [self.prepared, self.recalculated, self.output]
        return self.cli(command, *args, '--profile', self.profile, '--month', '2026-10', '--issue-date', '2026-10-15', '--due-date', '2026-10-31')

    def test_formula_dates_are_fixed_only_at_approved_cells_with_explicit_issue_date(self):
        self.volatile_date_template(with_print_area=True)

        prepared = self.dated_cli('prepare')

        self.assertEqual(prepared.returncode, 0, prepared.stderr)
        ws = load_workbook(self.prepared)['請求書 ']
        self.assertEqual(ws['G3'].value, datetime(2026, 10, 15))
        self.assertEqual(ws['B7'].value, datetime(2026, 10, 31))
        self.assertEqual(ws['B8'].value, '=TODAY()')
        self.assertEqual(str(ws.print_area), "'請求書 '!$A$1:$G$12")
        self.recalculate({'B8': '46310', 'G10': '129.6'})
        result = self.dated_cli('export')
        self.assertEqual(result.returncode, 0, result.stderr)
        final = load_workbook(self.output)['請求書 ']
        self.assertEqual(final['G3'].value, datetime(2026, 10, 15))
        self.assertEqual(final['B7'].value, datetime(2026, 10, 31))
        self.assertEqual(final['G10'].value, 129.6)
        self.assertFalse(any(c.data_type == 'f' for row in final for c in row))
        self.assertEqual(json.loads(result.stdout)['print_preview'], 'manual_required')

    def test_unconfigured_print_area_is_not_guessed_or_reported_as_verified(self):
        self.volatile_date_template(with_print_area=False)
        prepared = self.dated_cli('prepare')
        self.assertEqual(prepared.returncode, 0, prepared.stderr)
        self.assertFalse(load_workbook(self.prepared)['請求書 '].print_area)
        self.recalculate({'B8': '46310', 'G10': '129.6'})

        result = self.dated_cli('export')

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('print_areas', result.stderr)
        self.assertFalse(self.output.exists())

    def test_date_formula_override_requires_the_requested_explicit_date(self):
        self.volatile_date_template(with_print_area=True)

        result = self.cli('prepare', self.template, self.prepared, '--profile', self.profile, '--month', '2026-10', '--issue-date', '2026-10-15')

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('明示的な日付引数', result.stderr)
        self.assertFalse(self.prepared.exists())


    def test_export_copies_only_caches_and_keeps_original_styles_and_margins(self):
        self.prepare()
        self.recalculate()
        def changed_style(name, data):
            if name == 'xl/styles.xml':
                return data.replace(b'Calibri', b'DifferentFont')
            if name == 'xl/worksheets/sheet1.xml':
                return data.replace(b'left="0.75"', b'left="0.75033"').replace(b'header="0.5"', b'header="0.511811"')
            return data
        self.rewrite_zip(self.recalculated, changed_style)

        result = self.export()

        self.assertEqual(result.returncode, 0, result.stderr)
        with ZipFile(self.prepared) as src, ZipFile(self.output) as dst:
            self.assertEqual(src.read('xl/styles.xml'), dst.read('xl/styles.xml'))
        self.assertEqual(load_workbook(self.output).active.page_margins.left, 0.75)
        self.assertEqual(load_workbook(self.output).active.page_margins.header, 0.5)

    def test_profile_layout_changes_only_explicit_columns_and_empty_merge_cells(self):
        self.volatile_date_template(with_print_area=True)
        profile = json.loads(self.profile.read_text())
        profile['layout'] = {'請求書 ': {'column_widths': {'G': 18}, 'merge_ranges': ['B7:C7']}}
        self.profile.write_text(json.dumps(profile))

        result = self.dated_cli('prepare')

        self.assertEqual(result.returncode, 0, result.stderr)
        ws = load_workbook(self.prepared)['請求書 ']
        self.assertEqual(ws.column_dimensions['G'].width, 18)
        self.assertIn('B7:C7', ws.merged_cells)
        self.assertEqual(ws['B7'].value, datetime(2026, 10, 31))
        self.assertEqual(ws['B8'].value, '=TODAY()')

    def test_profile_merge_rejects_nonempty_neighbor_instead_of_losing_data(self):
        self.volatile_date_template(with_print_area=True)
        wb = load_workbook(self.template)
        wb['請求書 ']['C7'] = 'Keep this value'
        wb.save(self.template)
        profile = json.loads(self.profile.read_text())
        profile['layout'] = {'請求書 ': {'merge_ranges': ['B7:C7']}}
        self.profile.write_text(json.dumps(profile))

        result = self.dated_cli('prepare')

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('結合', result.stderr)
        self.assertFalse(self.prepared.exists())


    def test_profile_existing_merge_with_absolute_coordinates_is_not_duplicated(self):
        self.volatile_date_template(with_print_area=True)
        wb = load_workbook(self.template)
        wb['請求書 '].merge_cells('B7:C7')
        wb.save(self.template)
        profile = json.loads(self.profile.read_text())
        profile['layout'] = {'請求書 ': {'merge_ranges': ['$B$7:$C$7']}}
        self.profile.write_text(json.dumps(profile))

        result = self.dated_cli('prepare')

        self.assertEqual(result.returncode, 0, result.stderr)
        with ZipFile(self.prepared) as archive:
            xml = ET.fromstring(archive.read('xl/worksheets/sheet1.xml'))
            self.assertEqual(len(xml.findall('.//s:mergeCell', NS)), 1)

    def test_export_rejects_changed_date_epoch_even_if_read_cell_values_match(self):
        self.prepare()
        self.recalculate()
        def change_epoch(name, data):
            if name == 'xl/workbook.xml':
                return data.replace(b'<workbookPr', b'<workbookPr date1904="1"', 1)
            if name == 'xl/worksheets/sheet1.xml':
                # 入力日付の serial を変えて、読み取り時の日付は同じに保つ。
                return data.replace(b'46296', b'44834')
            return data
        self.rewrite_zip(self.recalculated, change_epoch)

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('日付基準', result.stderr)
        self.assertFalse(self.output.exists())


    def test_prepare_rejects_date_override_inside_shared_formula_group(self):
        self.volatile_date_template(with_print_area=True)
        wb = load_workbook(self.template)
        wb['請求書 ']['G4'] = '=TODAY()'
        wb['請求書 ']['G4'].number_format = 'yyyy-mm-dd'
        wb.save(self.template)
        def shared_dates(name, data):
            if name == 'xl/worksheets/sheet1.xml':
                xml = ET.fromstring(data)
                for node in xml.findall('.//s:c', NS):
                    if node.get('r') in ('G3', 'G4'):
                        formula = node.find('s:f', NS)
                        formula.set('t', 'shared')
                        formula.set('si', '0')
                        if node.get('r') == 'G3':
                            formula.set('ref', 'G3:G4')
                        else:
                            formula.text = None
                return ET.tostring(xml)
            return data
        self.rewrite_zip(self.template, shared_dates)

        result = self.dated_cli('prepare')

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('共有数式', result.stderr)
        self.assertFalse(self.prepared.exists())

    def test_prepare_removes_chain_references_after_date_formulas_are_fixed(self):
        self.volatile_date_template(with_print_area=True)
        def add_chain(name, data):
            if name == '[Content_Types].xml':
                return data.replace(b'</Types>', b'<Override PartName="/xl/calcChain.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.calcChain+xml"/></Types>')
            if name == 'xl/_rels/workbook.xml.rels':
                return data.replace(b'</Relationships>', b'<Relationship Id="chain" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/calcChain" Target="calcChain.xml"/></Relationships>')
            return data
        self.rewrite_zip(self.template, add_chain, {'xl/calcChain.xml': '<calcChain xmlns="%s"><c r="G3" i="1"/><c r="B7"/></calcChain>' % NS['s']})

        result = self.dated_cli('prepare')

        self.assertEqual(result.returncode, 0, result.stderr)
        with ZipFile(self.prepared) as archive:
            self.assertNotIn('xl/calcChain.xml', archive.namelist())
            self.assertNotIn(b'calcChain', archive.read('[Content_Types].xml'))
            self.assertNotIn(b'calcChain', archive.read('xl/_rels/workbook.xml.rels'))


    def change_defined_name_attributes(self, path, updates):
        def change(name, data):
            if name == 'xl/workbook.xml':
                xml = ET.fromstring(data)
                nodes = xml.findall('.//s:definedName', NS)
                self.assertTrue(nodes)
                for node in nodes:
                    for key, value in updates.items():
                        if value is None:
                            node.attrib.pop(key, None)
                        else:
                            node.set(key, value)
                return ET.tostring(xml)
            return data
        self.rewrite_zip(path, change)

    def test_export_accepts_default_false_defined_name_flags_added_by_recalculation(self):
        self.prepare()
        for false_value in ('false', '0'):
            with self.subTest(false_value=false_value):
                self.output.unlink(missing_ok=True)
                self.recalculate()
                self.change_defined_name_attributes(self.recalculated, dict.fromkeys(('function', 'hidden', 'vbProcedure'), false_value))

                result = self.export()

                self.assertEqual(result.returncode, 0, result.stderr)
                with ZipFile(self.prepared) as source, ZipFile(self.output) as output:
                    self.assertEqual(source.read('xl/workbook.xml'), output.read('xl/workbook.xml'))

    def test_export_accepts_explicit_false_defined_name_flags_becoming_unspecified(self):
        self.change_defined_name_attributes(self.template, {'function': 'false', 'hidden': '0', 'vbProcedure': 'false'})
        self.prepare()
        self.recalculate()
        self.change_defined_name_attributes(self.recalculated, dict.fromkeys(('function', 'hidden', 'vbProcedure'), None))

        result = self.export()

        self.assertEqual(result.returncode, 0, result.stderr)

    def test_export_still_rejects_true_flags_and_other_defined_name_attribute_changes(self):
        self.prepare()
        for key, value in [('function', 'true'), ('hidden', '1'), ('vbProcedure', 'true'), ('name', 'ChangedName'), ('comment', 'false')]:
            with self.subTest(attribute=key):
                self.recalculate()
                self.change_defined_name_attributes(self.recalculated, {key: value})

                result = self.export()

                self.assertNotEqual(result.returncode, 0)
                self.assertIn('入力値・数式', result.stderr)
                self.assertFalse(self.output.exists())

    def test_export_still_rejects_defined_name_scope_changes(self):
        wb = load_workbook(self.template)
        wb.create_sheet('Other')
        wb.save(self.template)
        self.prepare()
        self.recalculate()
        self.change_defined_name_attributes(self.recalculated, {'localSheetId': '1'})

        result = self.export()

        self.assertNotEqual(result.returncode, 0)
        self.assertIn('入力値・数式', result.stderr)
        self.assertFalse(self.output.exists())


if __name__ == '__main__':
    unittest.main()
